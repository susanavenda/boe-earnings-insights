#!/usr/bin/env python3
"""Build Group 9 Assignment 3 report (HTML, DOCX, PDF) and slides (HTML, PDF)."""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

NAVY = RGBColor(0x1B, 0x3A, 0x5C)
GREY = RGBColor(0x6B, 0x6B, 0x68)
INK = RGBColor(0x1A, 0x1A, 0x18)

# Prose counted toward the 1,500-word limit. Tables, captions, title block,
# and the code snapshot are excluded, matching the Assignment 1 convention.
S1 = [
    "The Prudential Regulation Authority supervises around 1,500 firms, mainly through structured returns: data on a defined schedule, checked against defined fields. Quarterly results sit outside that pipeline. They mix a curated management presentation with unscripted analyst Q&A, and reading them properly takes both financial and technical skill, so any signal in them is not used systematically. The question is not whether a call can be summarised. It is whether earnings announcements, analyst Q&A, and other public material carry a leading signal about a firm’s prudential condition that the reported metrics alone do not.",
    "Leading and alone are the words that do the work. A cheerful tone that never answers the question is not the same evidence as a number that moved. This project answers with a repeatable pipeline, not a rating of a firm or of a person. Public PDFs and Excel packs go in. What comes out is a pack a desk can check: the theme, the tone, and whether the answer was direct, supplied the metric, or changed the subject. Re-running it next quarter takes about 45 to 90 minutes. A null — no early-warning edge from the call — is a valid result.",
]

S2 = [
    "The first decision was what to measure, because a sentiment label only describes the conversation. Beside tone, the pipeline scores the answer itself: was it direct, did it supply the metric that was asked for, and was the topic substituted. In one worked case an analyst asked for the 90-day past-due rate on the UK cards book, this quarter against the last. Management said they were comfortable with the overall credit picture and that the portfolio remained robust. The tone is positive. The record is also evasive: directness is low, the requested number is absent, and UK cards has been replaced by “overall credit”. If directness on one topic falls while the tone stays cheerful, that is a warning a supervisor can use. Sentiment alone is not enough.",
    "The second decision was to fix the categories before anyone read a transcript. Eighteen prudential categories were taken from published PRA and Basel material and linked to reporting families a supervisor already uses, including LCR, NSFR and ALMM for liquidity. Capital is split into the current position and forward-looking stress planning. Profitability keeps an interest-rate element for IRRBB. Mixed and untagged rows stay unmapped until a person separates them. Built this way, the map is a test. Built from the transcripts, it would succeed by design. Where it fails is informative: if liquidity is first-order for the PRA and the call barely mentions it, that absence is a finding. A later eight-way filing of the same map groups quotes. It does not fire an alert.",
    "Scope stayed on HSBC and Barclays, both UK-incorporated and under full PRA oversight. The plan named 20 transcripts. The finished UK set is 137 files, 136 of them results calls. One HSBC equity-analyst meeting is a different event and is excluded rather than averaged in. Peer comparison uses the 62 quarters in which both banks published, paired by quarter rather than by complete year, which keeps 2010 and 2011. HSBC reported twice a year until 2008. Barclays does not appear until the second quarter of 2010. Gaps stay gaps. Nothing is interpolated or filled with zero. The four metrics stayed locked: total income, operating costs, credit impairment or ECL, and CET1. Credit Suisse 2020–2022 was built as an out-of-sample third bank and was held out of the pitch. It is in this report. It does not enter the HSBC–Barclays peer gap.",
    "Preparation is six steps from PDF to a scored pair, because an unnamed turn would poison question-versus-answer sentiment. pdfplumber extracts the text. Each bank has its own speaker rule: HSBC puts name and role on one line, Barclays puts name and firm on their own line, and the Credit Suisse pattern is tried first so it does not steal other speakers. That yields 4,188 speaker turns, and every analyst turn is named. Those turns become 1,213 question-and-answer pairs. In 192 cases an executive reply sat inside the analyst’s turn and had to be split. HSBC’s “interim” and Barclays’ “Q2” keep their original labels and also receive a shared calendar period. The store is SQLite, which the notebook rebuilds. On Colab, Stage 0 clones the public repository and runs the scripts from that clone. Among questions about the four metrics, 600 answers address the metric and 218 do not, coverage of 73%. The record keeps its limits: 108 empty answers and 369 rows with no publication date.",
    "Two events were chosen as temporal tests, and they do not prove the same thing. Late 2019 into early 2020 asks whether directness on credit and provisioning fell before the lockdowns. The 2022–23 stress asks whether directness on liquidity dropped in the months before SVB and Credit Suisse failed. COVID was an external shock, so a move in the call is association, not a prediction. The peer window for a matched gap starts in 2012. A Yahoo Finance check covers recent headlines only, not the full twenty years. Time, peer, and pack versus narrative are the baselines. Both a divergence and an agreement count as findings.",
    "Topics use BERTopic rather than LDA because short financial questions need dense embeddings. On this corpus BERTopic coherence is 0.46 against LDA at 0.40, with five live topics, so LDA stays the check. Topic numbers move if the model is refit, so a pack cites the quote, not a permanent topic id. Sentiment was scored blind against one human coder on 90 pairs. Whole-turn FinBERT agrees 60% of the time, Cohen’s kappa 0.28, macro-F1 0.52. The Loughran–McDonald lexicon agrees 42%, kappa 0.02, which is chance once the shared habit of choosing neutral is removed. Kappa is the figure we stand behind, not 60% on its own. FinBERT labels about 80% of questions neutral and misses most human negatives, which is the hedged language supervisors care about. The two sentiment methods agree with each other only about 58% of the time. Sentence scoring and an LLM coder were also run. Neither met the 70% bar set before coding began, and the bar was not moved.",
    "A light fine-tune of FinBERT was trained and not promoted. A gain on labels the model family helped to create would be circular, so it is not a reason to change the head. The gate stays shut until the second coder’s blind pass is in, and only a win against plain FinBERT on those human labels would open it. Scripts score the human files. They do not rewrite them. The pack itself is rules, not one risk score, and alert, watch, and null are not a PRA rating. Alert needs a credit narrative that disagrees with the pack, or a matched peer gap worse than −0.10 where neither side is entirely FinBERT-neutral. Watch is a soft net below −0.08 without that contradiction. Null is flat tone: no early-warning edge.",
    "The proof cases are where the pitch becomes a product. HSBC half-year 2025 is watch: FinBERT net −0.13 on eight analyst turns. The matched Barclays print is entirely FinBERT-neutral, so the gap of about −0.13 is shown and refused as an alert. Barclays half-year 2026 is the same shape of watch in the other direction. Credit Suisse’s fourth quarter of 2022, the failure window held out of the pitch, is null: net +0.08 on 20 turns and no FinBERT negatives. The lexicon flags loss words inside some of those neutral turns. That is disclosed on the note. It does not flip the rule. Across 2020–2022 the Credit Suisse tone does not worsen into the collapse, so the pipeline is not read as having predicted a failure it did not see.",
    "Two constraints sit on the method, taken from the Bank’s own trusted-AI tests rather than from an internal checklist. Directness is a triage signal. It says where to look. It is not an assessment of a firm or of an individual, and no conclusion follows without a person. The regulatory regime also changed under the data. The PRA did not exist before 2013, so a category built from today’s PRA sources, laid on a 2006 transcript, is anachronistic by construction. The run is provisional for that reason. On 212 bank-quarter-metric overlaps, narrative direction and the pack agree 26% of the time, and credit impairment agrees on 18% of 44 rows. Many of those splits are a flat narrative against a pack that moved, which is the neutral bias again, so they are a lead for a person, not an automatic alert.",
]

S3 = [
    "The strongest result is a negative one. The two sentiment methods agree with each other only about 58% of the time. FinBERT calls about 80% of questions neutral. Against a human coder it reaches kappa 0.28, under the 70% bar set before the work started, and it misses the hedged negatives a supervisor would want to see. The behavioural record — directness, whether the metric was supplied, whether the topic was substituted — is what the pack adds on top of tone. HSBC half-year 2025 is watch, not a rating of the firm, and the peer alert is refused because the matched Barclays print is model-flat. Credit Suisse’s failure quarter, held out of the pitch, is null.",
    "For the Bank, use the pack as triage when filings land. A person accepts or rejects watch, alert, or null. Do not promote the fine-tune until the second coder’s blind labels are in, and only if it beats plain FinBERT there. Do not fill gaps in the peer history, and do not treat a neutral score as “no issue”. Next is that second coder, a retag of metric hits so an agree column is evidenced rather than empty, and a clear note wherever a current PRA category is laid on a year before 2013.",
]

EXEC = [
    "Group 9 asked whether earnings Q&A carries a leading signal that reported metrics alone do not. The pipeline scores tone and, separately, whether the answer was direct, supplied the number, and stayed on the topic asked. The eighteen risk categories were taken from PRA material before any transcript was read. FinBERT agrees with one human coder 60% of the time, kappa 0.28, under a 70% bar set in advance. HSBC half-year 2025 is watch. Credit Suisse’s 2022 failure quarter, held out of the pitch, is null.",
]


def words(paragraphs: list[str]) -> int:
    return sum(len(p.split()) for p in paragraphs)


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, *, size, bold=False, color=None, name="Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_p(doc, text, *, size=10.5, bold=False, color=None, space_before=0, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    r.italic = italic
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run_font(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "1B3A5C")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(val)
            set_run_font(r, size=10, color=INK)
            if ri % 2 == 1:
                shade(cell, "F4F7FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    add_p(doc, "TECHNICAL REPORT", size=10, bold=True, color=GREY, space_after=2)
    add_p(doc, "Assignment 3: Final report", size=18, bold=True, color=NAVY, space_after=6)
    add_p(
        doc,
        "Employer project with the Bank of England (Group 9). Reports the pipeline from earnings Q&A to an alert, watch, or null pack, the decisions behind it, and what a supervisor should do with the result.",
        size=11,
        color=RGBColor(0x3C, 0x3C, 0x3A),
        space_after=4,
    )
    add_p(
        doc,
        "Cambridge Data Science Career Accelerator · Group 9 · 2026",
        size=10,
        color=NAVY,
        space_after=2,
    )
    add_p(
        doc,
        "Banks: HSBC · Barclays    Methods: BERTopic · LDA · FinBERT · LDSA · rules protocol    Store: data/boe.sqlite",
        size=9,
        color=GREY,
        space_after=10,
    )

    add_p(doc, "Executive summary", size=16, bold=True, color=NAVY, space_before=6, space_after=2)
    add_p(
        doc,
        "Front matter. Not counted in the 1,500-word body, following the Assignment 1 convention. Counted prose is Sections 1–3.",
        size=9,
        italic=True,
        color=GREY,
        space_after=6,
    )
    for para in EXEC:
        add_p(doc, para, size=10.5, space_after=8)

    add_p(doc, "1. Background and context", size=16, bold=True, color=NAVY, space_before=8, space_after=6)
    for para in S1:
        add_p(doc, para, size=10.5, space_after=8)
    add_table(
        doc,
        ["Decision", "What was locked"],
        [
            ["Banks", "HSBC and Barclays. UK-incorporated, full PRA oversight."],
            ["Question", "Does Q&A carry a signal the four reported lines miss?"],
            ["Output", "Alert, watch, or null, with quotes and sample size. Not a firm rating."],
            ["Out of scope", "Video, private returns, training a model from scratch, extra banks."],
        ],
    )
    add_p(doc, "Table 1. Scope carried from the plan into the finished pipeline.", size=9, color=GREY, space_before=2)

    add_p(doc, "2. Project development", size=16, bold=True, color=NAVY, space_before=10, space_after=6)
    for para in S2:
        add_p(doc, para, size=10.5, space_after=8)
    add_table(
        doc,
        ["Choice", "Alternative considered", "Why this one"],
        [
            ["BERTopic", "Fixed taxonomy only, or LDA as the live model", "Themes come from the Q&A. LDA is the coherence check (0.46 vs 0.40)."],
            ["Whole-turn FinBERT", "Sentence FinBERT, lexicon, or LLM as the scorer", "Best of the tested units on answers. None met 70%."],
            ["Rules protocol", "A single risk score", "A desk can re-run and refuse a bad peer gap."],
            ["Human gold gate", "Promote the fine-tune on silver labels", "Silver lift was circular (0.43 to 0.64). Head stays unpromoted."],
        ],
    )
    add_p(doc, "Table 2. Model choices against the alternatives that were actually run.", size=9, color=GREY, space_before=2)
    add_p(doc, "Protocol conditions, as implemented.", size=10.5, bold=True, color=NAVY, space_before=8, space_after=4)
    add_p(
        doc,
        "A1 alert  — topic negative share rises AND credit-impairment narrative disagrees with the pack.\n"
        "A2 alert  — matched peer gap < −0.10, both n ≥ 3, peer side not 100% FinBERT-neutral.\n"
        "A3 watch  — FinBERT net < −0.08 and the pack is not contradicted.\n"
        "N1 null   — tone flat and pack and narrative agree. Report no early-warning edge.",
        size=10,
        space_after=4,
    )
    add_p(doc, "Code snapshot. Rules in scripts/build_supervisory_episode.py, published by scripts/generate_pra_note.py.", size=9, color=GREY)

    add_p(doc, "3. Results", size=16, bold=True, color=NAVY, space_before=10, space_after=6)
    for para in S3:
        add_p(doc, para, size=10.5, space_after=8)
    add_table(
        doc,
        ["Episode", "FinBERT net", "n", "Verdict"],
        [
            ["HSBC 2025-H1", "−0.13", "8", "Watch (A3). Peer alert refused."],
            ["Barclays 2026-H1", "−0.11", "7", "Watch (A3). Peer alert refused."],
            ["Credit Suisse 2022-Q4", "+0.08", "20", "Null (N1). Out of sample."],
        ],
    )
    add_p(doc, "Table 3. Reviewed protocol cases. Source: PRA notes generated 27 September 2026.", size=9, color=GREY, space_before=2)
    add_table(
        doc,
        ["Check", "Result"],
        [
            ["FinBERT vs human gold", "60% raw, kappa 0.28, macro-F1 0.52. Bar of 70% not met."],
            ["Eight-way dual code", "Machine vs machine 35%. Human vs machine 50% (n=60)."],
            ["Sentence FinBERT / LLM", "Sentence kappa 0.13. LLM negative recall on questions 69% vs FinBERT 31%."],
            ["Pack vs narrative", "Agree on 26% of 212 overlaps. Credit impairment 18% of 44."],
            ["Fine-tune", "Not promoted. active_model_id stays empty."],
        ],
    )
    add_p(doc, "Table 4. Validation the verdict depends on.", size=9, color=GREY, space_before=2)

    add_p(doc, "Reproducibility", size=16, bold=True, color=NAVY, space_before=10, space_after=6)
    add_p(
        doc,
        "Notebook: notebooks/boe_earnings_insights.ipynb. On Colab, Stage 0 clones https://github.com/susanavenda/boe-earnings-insights and runs scripts/ from that clone. CSV mirrors of the tables are in data/processed/. Team: Aidan, Alfred, Bupathi, Debanjan, Rafael, Susana Venda, Taz.",
        size=10.5,
        space_after=8,
    )
    n1, n2, n3 = words(S1), words(S2), words(S3)
    add_p(
        doc,
        f"Word count: {n1 + n2 + n3} words in Sections 1–3 (Section 1: {n1}; Section 2: {n2}; Section 3: {n3}). Tables, captions, the code snapshot, the executive summary, and this line are excluded.",
        size=9,
        color=GREY,
    )
    doc.save(path)


REPORT_CSS = """
@page { size: A4; margin: 16mm 16mm 18mm 16mm; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #1a1a18;
  font-family: Arial, Helvetica, sans-serif; font-size: 10.5pt; line-height: 1.35; }
.kicker { font-size: 10pt; font-weight: 700; color: #6b6b68; letter-spacing: 0.04em; margin: 0 0 4px; }
h1 { font-size: 18pt; line-height: 1.2; color: #1b3a5c; margin: 0 0 8px; font-weight: 700; }
h2 { font-size: 16pt; color: #1b3a5c; margin: 16px 0 6px; font-weight: 700; }
.deck { font-size: 11pt; color: #3c3c3a; margin: 0 0 6px; }
.meta { font-size: 10pt; color: #1b3a5c; margin: 0 0 2px; }
.sub { font-size: 9pt; color: #6b6b68; margin: 0 0 10px; }
.note { font-size: 9pt; color: #6b6b68; font-style: italic; margin: 0 0 8px; }
p { margin: 0 0 8px; }
table { width: 100%; border-collapse: collapse; margin: 6px 0 2px; }
th, td { border: 1px solid #d9e2ec; padding: 4px 6px; text-align: left; vertical-align: top; font-size: 10pt; }
th { background: #1b3a5c; color: #fff; font-weight: 700; }
tr:nth-child(even) td { background: #f4f7fa; }
caption { caption-side: bottom; text-align: left; font-size: 9pt; color: #6b6b68; margin-top: 2px; padding-top: 2px; }
pre { font-family: "Courier New", Courier, monospace; font-size: 10pt; background: #f4f7fa;
  border-left: 4px solid #1b3a5c; padding: 8px 10px; white-space: pre-wrap; margin: 4px 0; }
"""

SLIDE_CSS = """
@page { size: 13.333in 7.5in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #0b1f33;
  font-family: Arial, Helvetica, sans-serif; font-size: 14pt; line-height: 1.28; }
.slide { width: 13.333in; height: 7.5in; padding: 0.38in 0.48in 0.42in;
  page-break-after: always; break-after: page; position: relative; overflow: hidden; background: #fff; }
.slide:last-child { page-break-after: auto; break-after: auto; }
.kicker { font-size: 14pt; color: #486581; letter-spacing: 0.04em; text-transform: uppercase;
  font-weight: 700; margin: 0 0 0.08in; }
h1 { font-size: 28pt; line-height: 1.15; margin: 0 0 0.18in; font-weight: 700; }
h2 { font-size: 22pt; line-height: 1.15; margin: 0 0 0.16in; font-weight: 700; }
h3 { font-size: 14pt; margin: 0 0 0.08in; font-weight: 700; color: #1a3d5c; }
p { margin: 0 0 0.1in; }
.footer { position: absolute; left: 0.48in; right: 0.48in; bottom: 0.18in;
  display: flex; justify-content: space-between; font-size: 14pt; color: #627d98;
  border-top: 2px solid #c4a35a; padding-top: 0.08in; }
.stats { display: flex; gap: 0.16in; margin: 0.12in 0 0.16in; }
.stat { flex: 1; background: #f4f7fa; border: 1px solid #d9e2ec; padding: 0.12in 0.14in; text-align: center; }
.stat .v { font-size: 22pt; font-weight: 700; color: #0b1f33; }
.stat .l { font-size: 14pt; color: #486581; margin-top: 0.04in; }
.stat.warn .v { color: #6b4e12; }
.stat.danger .v { color: #7a1f12; }
.callout { background: #fff8eb; border-left: 6px solid #c4a35a; padding: 0.12in 0.16in; margin: 0.1in 0 0.14in; }
.callout.info { background: #eef4f8; border-left-color: #1a3d5c; }
.callout.ok { background: #e8f3ea; border-left-color: #1e4d2b; }
.callout.bad { background: #f8ece8; border-left-color: #7a1f12; }
.callout .t { font-weight: 700; margin-bottom: 0.04in; }
table { width: 100%; border-collapse: collapse; margin: 0.08in 0 0.12in; }
th, td { border: 1px solid #d9e2ec; padding: 0.06in 0.08in; text-align: left; vertical-align: top; font-size: 14pt; }
th { background: #0b1f33; color: #fff; font-weight: 700; }
tr:nth-child(even) td { background: #f7f9fb; }
.cols { display: flex; gap: 0.22in; }
.cols > div { flex: 1; }
.muted { color: #486581; }
.quote { background: #f7f9fb; border-left: 5px solid #1a3d5c; padding: 0.1in 0.14in; margin-top: 0.08in; }
"""


def paras(items: list[str]) -> str:
    return "\n".join(f"<p>{p}</p>" for p in items)


def build_report_html(path: Path) -> None:
    n1, n2, n3 = words(S1), words(S2), words(S3)
    html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"/>
<title>Group 9 · CAM EP Assignment 3 · Report</title>
<style>{REPORT_CSS}</style></head><body>
<p class="kicker">TECHNICAL REPORT</p>
<h1>Assignment 3: Final report</h1>
<p class="deck">Employer project with the Bank of England (Group 9). Reports the pipeline from earnings Q&amp;A to an alert, watch, or null pack, the decisions behind it, and what a supervisor should do with the result.</p>
<p class="meta">Cambridge Data Science Career Accelerator · Group 9 · 2026</p>
<p class="sub">Banks: HSBC · Barclays &nbsp;&nbsp; Methods: BERTopic · LDA · FinBERT · LDSA · rules protocol &nbsp;&nbsp; Store: data/boe.sqlite</p>
<h2>Executive summary</h2>
<p class="note">Front matter. Not counted in the 1,500-word body, following the Assignment 1 convention. Counted prose is Sections 1–3.</p>
{paras(EXEC)}
<h2>1. Background and context</h2>
{paras(S1)}
<table>
<tr><th>Decision</th><th>What was locked</th></tr>
<tr><td>Banks</td><td>HSBC and Barclays. UK-incorporated, full PRA oversight.</td></tr>
<tr><td>Question</td><td>Does Q&amp;A carry a signal the four reported lines miss?</td></tr>
<tr><td>Output</td><td>Alert, watch, or null, with quotes and sample size. Not a firm rating.</td></tr>
<tr><td>Out of scope</td><td>Video, private returns, training a model from scratch, extra banks.</td></tr>
</table>
<p class="sub">Table 1. Scope carried from the plan into the finished pipeline.</p>
<h2>2. Project development</h2>
{paras(S2)}
<table>
<tr><th>Choice</th><th>Alternative considered</th><th>Why this one</th></tr>
<tr><td>BERTopic</td><td>Fixed taxonomy only, or LDA as the live model</td><td>Themes come from the Q&amp;A. LDA is the coherence check (0.46 vs 0.40).</td></tr>
<tr><td>Whole-turn FinBERT</td><td>Sentence FinBERT, lexicon, or LLM as the scorer</td><td>Best of the tested units on answers. None met 70%.</td></tr>
<tr><td>Rules protocol</td><td>A single risk score</td><td>A desk can re-run and refuse a bad peer gap.</td></tr>
<tr><td>Human gold gate</td><td>Promote the fine-tune on silver labels</td><td>Silver lift was circular (0.43 to 0.64). Head stays unpromoted.</td></tr>
</table>
<p class="sub">Table 2. Model choices against the alternatives that were actually run.</p>
<p><strong>Protocol conditions, as implemented.</strong></p>
<pre>A1 alert  — topic negative share rises AND credit-impairment narrative disagrees with the pack.
A2 alert  — matched peer gap &lt; −0.10, both n ≥ 3, peer side not 100% FinBERT-neutral.
A3 watch  — FinBERT net &lt; −0.08 and the pack is not contradicted.
N1 null   — tone flat and pack and narrative agree. Report no early-warning edge.</pre>
<p class="sub">Code snapshot. Rules in scripts/build_supervisory_episode.py, published by scripts/generate_pra_note.py.</p>
<h2>3. Results</h2>
{paras(S3)}
<table>
<tr><th>Episode</th><th>FinBERT net</th><th>n</th><th>Verdict</th></tr>
<tr><td>HSBC 2025-H1</td><td>−0.13</td><td>8</td><td>Watch (A3). Peer alert refused.</td></tr>
<tr><td>Barclays 2026-H1</td><td>−0.11</td><td>7</td><td>Watch (A3). Peer alert refused.</td></tr>
<tr><td>Credit Suisse 2022-Q4</td><td>+0.08</td><td>20</td><td>Null (N1). Out of sample.</td></tr>
</table>
<p class="sub">Table 3. Reviewed protocol cases. Source: PRA notes generated 27 September 2026.</p>
<table>
<tr><th>Check</th><th>Result</th></tr>
<tr><td>FinBERT vs human gold</td><td>60% raw, kappa 0.28, macro-F1 0.52. Bar of 70% not met.</td></tr>
<tr><td>Eight-way dual code</td><td>Machine vs machine 35%. Human vs machine 50% (n=60).</td></tr>
<tr><td>Sentence FinBERT / LLM</td><td>Sentence kappa 0.13. LLM negative recall on questions 69% vs FinBERT 31%.</td></tr>
<tr><td>Pack vs narrative</td><td>Agree on 26% of 212 overlaps. Credit impairment 18% of 44.</td></tr>
<tr><td>Fine-tune</td><td>Not promoted. active_model_id stays empty.</td></tr>
</table>
<p class="sub">Table 4. Validation the verdict depends on.</p>
<h2>Reproducibility</h2>
<p>Notebook: notebooks/boe_earnings_insights.ipynb. On Colab, Stage 0 clones https://github.com/susanavenda/boe-earnings-insights and runs scripts/ from that clone. CSV mirrors of the tables are in data/processed/. Team: Aidan, Alfred, Bupathi, Debanjan, Rafael, Susana Venda, Taz.</p>
<p class="sub">Word count: {n1 + n2 + n3} words in Sections 1–3 (Section 1: {n1}; Section 2: {n2}; Section 3: {n3}). Tables, captions, the code snapshot, the executive summary, and this line are excluded.</p>
</body></html>
"""
    path.write_text(html, encoding="utf-8")


def slide(kicker, title, body, n, total=12, h="h2"):
    return f"""<section class="slide">
  <div class="kicker">{kicker}</div>
  <{h}>{title}</{h}>
  {body}
  <div class="footer"><span>Group 9 · BoE employer project</span><span>{n} / {total}</span></div>
</section>
"""


def build_slides_html(path: Path) -> None:
    slides = []
    slides.append(slide(
        "Group 9 · Bank of England PRA · Assignment 3 · 15 minutes",
        "Supervisory earnings pipeline",
        """
<div class="callout"><div class="t">The question</div>
Not whether a call can be summarised. Whether earnings Q&amp;A carries a <strong>leading</strong> signal that reported metrics <strong>alone</strong> do not.</div>
<div class="stats">
  <div class="stat"><div class="v">1,500</div><div class="l">Firms the PRA supervises</div></div>
  <div class="stat"><div class="v">137</div><div class="l">UK results transcripts</div></div>
  <div class="stat warn"><div class="v">0.28</div><div class="l">FinBERT kappa, not 60%</div></div>
  <div class="stat warn"><div class="v">WATCH</div><div class="l">HSBC H1 2025</div></div>
</div>
<p>These slides are the story told in the pitch, closed with the cases held for this report. The notebook is the technical walkthrough.</p>
""",
        1, h="h1"))
    slides.append(slide(
        "Background · Speaks: Taz",
        "Structured returns are the pipeline. The call is not.",
        """
<div class="cols">
  <div>
    <h3>What the PRA already runs</h3>
    <p>About 1,500 firms. Returns on a defined schedule, checked against defined fields.</p>
    <h3>What stays outside</h3>
    <p>Quarterly results: a curated presentation plus unscripted Q&amp;A. Expensive to read properly, so the signal is not used systematically.</p>
  </div>
  <div>
    <div class="callout info"><div class="t">Two words</div>
    Leading: did the call move before the stress, not after. Alone: information the four reported lines do not already contain.</div>
  </div>
</div>
<p><strong>What we measure besides tone:</strong> was the answer direct, did it supply the metric, or did it change the subject. Positive and evasive can be the same sentence.</p>
""",
        2))
    slides.append(slide(
        "Background · Speaks: Bupathi",
        "Positive tone, evasive answer",
        """
<div class="quote">Analyst: the 90-day past-due rate on the UK cards book, this quarter against the last. Management: comfortable with the overall credit picture; the portfolio remains robust.</div>
<table>
<tr><th>What a tone model writes</th><th>What we also write</th></tr>
<tr><td>Positive. The reply sounds confident.</td><td>Directness low. The number asked for was not given.</td></tr>
<tr><td>One label for the whole answer.</td><td>Topic substituted: UK cards became “overall credit”.</td></tr>
</table>
<div class="callout"><div class="t">Why it matters</div>
If directness on one topic falls while the tone stays cheerful, across 62 paired quarters, that is the warning. Sentiment alone is not enough.</div>
""",
        3))
    slides.append(slide(
        "Development · Speaks: Susana",
        "Same chain as the pitch. Now it is the product.",
        """
<table>
<tr><th style="width:18%">Step</th><th>What happens</th><th style="width:28%">Owner</th></tr>
<tr><td>1 Ingest</td><td>PDFs and Excel. Files stay inputs.</td><td>Susana</td></tr>
<tr><td>2 Prepare</td><td>Speaker split, then analyst Q&amp;A pairs.</td><td>Susana</td></tr>
<tr><td>3 Model</td><td>BERTopic, FinBERT, lexicon, behaviour, eight-way map.</td><td>Debanjan · Alfred · Bupathi · Aidan</td></tr>
<tr><td>4 Compare</td><td>Time, matched peer, pack versus Q&amp;A.</td><td>Susana · Aidan</td></tr>
<tr><td>5 Decide</td><td>Rules to alert, watch, or null.</td><td>Aidan · Bupathi</td></tr>
<tr><td>6 Publish</td><td>PRA note. Notebook is the factory.</td><td>Rafael · Susana</td></tr>
</table>
<div class="callout"><div class="t">Re-run</div>
45–90 minutes when new IR files land. Recalibration is a human gate, not a job that retrains itself.</div>
""",
        4))
    slides.append(slide(
        "Development · data · Speaks: Susana",
        "137 files. Gaps left as gaps.",
        """
<table>
<tr><th>Fact from the build</th><th>Figure</th></tr>
<tr><td>UK transcripts. One equity-analyst meeting excluded.</td><td>137 files, 136 results calls</td></tr>
<tr><td>Quarters where both banks published. Paired by quarter, not by full year.</td><td>62 quarters</td></tr>
<tr><td>Speaker turns, then Q&amp;A pairs. 192 replies sat inside an analyst turn.</td><td>4,188 turns · 1,213 pairs</td></tr>
<tr><td>Answers that address the four metrics, among questions about them.</td><td>600 of 818 · 73%</td></tr>
<tr><td>Empty answers, and rows with no publication date. Kept, not patched.</td><td>108 · 369</td></tr>
</table>
<p class="muted">HSBC is twice a year until 2008. Barclays starts in 2010 Q2. Nothing is interpolated or filled with zero. Credit Suisse was built and held out of the pitch. It is in this report, and it is not in the peer gap.</p>
""",
        5))
    slides.append(slide(
        "Development · methods · Speaks: Aidan · Debanjan · Alfred",
        "Categories first. Then the models.",
        """
<table>
<tr><th style="width:28%">We kept</th><th>Why this order</th></tr>
<tr><td>18 PRA categories, before any transcript</td><td>A map built from the calls would succeed by design. Liquidity barely mentioned is a finding.</td></tr>
<tr><td>BERTopic, LDA as the check</td><td>Coherence 0.46 versus 0.40. Five topics. Ids are not permanent alert keys.</td></tr>
<tr><td>Whole-turn FinBERT</td><td>Lexicon kappa 0.02, chance. Sentence scoring kappa 0.13. Neither replaces the whole answer.</td></tr>
<tr><td>Behaviour, not a fourth sentiment class</td><td>Direct, metric supplied, topic substituted. Avoidance is behaviour.</td></tr>
</table>
<div class="callout"><div class="t">Fine-tune stays off</div>
A gain on labels the model helped to create is not a promotion. The second coder is still the gate. We promote only if the fine-tune beats plain FinBERT on human labels.</div>
""",
        6))
    slides.append(slide(
        "Development · validation · Speaks: Aidan · Alfred",
        "The 70% bar was not met",
        """
<div class="stats">
  <div class="stat warn"><div class="v">0.28</div><div class="l">FinBERT kappa</div></div>
  <div class="stat"><div class="v">0.02</div><div class="l">Lexicon kappa. Chance.</div></div>
  <div class="stat"><div class="v">80%</div><div class="l">Questions called neutral</div></div>
  <div class="stat danger"><div class="v">70%</div><div class="l">Bar set first. Not met.</div></div>
</div>
<div class="cols">
  <div>
    <h3>Say kappa, not 60%</h3>
    <p>Both methods reach for neutral, so raw agreement flatters them. FinBERT does something. The lexicon does not. They agree with each other about 58% of the time.</p>
  </div>
  <div>
    <h3>What we did not move</h3>
    <p>The 70% bar was set before coding. We published the miss. Directness is triage: where to look, not a judgement of a firm or a person.</p>
  </div>
</div>
<div class="callout bad"><div class="t">Anachronism</div>
The PRA did not exist before 2013. A category from today’s PRA sources, laid on a 2006 transcript, is anachronistic. The run is provisional for that reason.</div>
""",
        7))
    slides.append(slide(
        "Development · episode · Speaks: Rafael · Bupathi",
        "HSBC 2025 is watch. The peer gap is not an alert.",
        """
<div class="stats">
  <div class="stat danger"><div class="v">−0.13</div><div class="l">FinBERT net</div></div>
  <div class="stat"><div class="v">8</div><div class="l">Analyst turns</div></div>
  <div class="stat warn"><div class="v">n/a</div><div class="l">Metric hits in the call</div></div>
  <div class="stat warn"><div class="v">WATCH</div><div class="l">Rule A3</div></div>
</div>
<div class="cols">
  <div>
    <h3>Why alert did not fire</h3>
    <p>Matched Barclays on the same half-year is entirely FinBERT-neutral (n=7). The gap is about −0.13. Rule A2 refuses it. Absence of a metric hit is not recorded as agreement.</p>
    <div class="quote">The pack shows the quote and the sample size. It does not rate HSBC.</div>
  </div>
  <div>
    <h3>Held out of the pitch, in this report</h3>
    <table>
    <tr><th>Case</th><th>Verdict</th></tr>
    <tr><td>Barclays 2026-H1, net −0.11, n=7</td><td>Watch. Peer refused.</td></tr>
    <tr><td>Credit Suisse 2022-Q4, net +0.08, n=20</td><td>Null. Not a predicted collapse.</td></tr>
    </table>
  </div>
</div>
""",
        8))
    slides.append(slide(
        "Results · Speaks: Rafael · Taz",
        "What the finished run shows",
        """
<table>
<tr><th>Finding</th><th>Number</th><th>Do not claim</th></tr>
<tr><td>Q&amp;A is not a copy of the pack</td><td>Agree on 26% of 212 overlaps</td><td>That disagreement is automatically an alert</td></tr>
<tr><td>Credit narrative is the weak join</td><td>18% agree, 44 rows</td><td>That ECL is mis-stated</td></tr>
<tr><td>Soft, without a bad peer, is watch</td><td>HSBC −0.13, rule A3</td><td>That the firm is weaker</td></tr>
<tr><td>A known failure can score null</td><td>CS Q4 2022, FinBERT +0.08</td><td>That the pipeline saw the collapse coming</td></tr>
</table>
<p class="muted">BERTopic coherence 0.46, LDA 0.40. Both stay. Only BERTopic names the live themes, and even those names are not used as permanent alert keys.</p>
""",
        9))
    slides.append(slide(
        "Results · recommendations · Speaks: Taz",
        "What the Bank should do with this",
        """
<div class="cols">
  <div>
    <h3>Use now</h3>
    <p>Run the pack when filings land. A person accepts or rejects watch, alert, or null. Keep the peer-quality gate. File a null in the same way as a watch.</p>
    <h3>Do not do</h3>
    <p>Do not promote the fine-tuned head. Do not add banks to make the peer gap look busier. Do not treat neutral as “no issue”.</p>
  </div>
  <div>
    <div class="callout info"><div class="t">Next analysis</div>
    Second coder on the 90-pair gold. Retag the four metrics inside the episode so the agree column is evidenced rather than n/a. Only then reopen promotion.</div>
    <div class="callout"><div class="t">Limit</div>
    Eight turns is a print, not a population. Hedged answers still collapse toward neutral. That limit is on the note.</div>
  </div>
</div>
""",
        10))
    slides.append(slide(
        "Technical walkthrough · Speaks: Susana",
        "These slides are the story. The notebook is the proof.",
        """
<div class="cols">
  <div>
    <h3>Open the notebook, not every cell</h3>
    <table>
    <tr><th>Stage</th><th>Show</th></tr>
    <tr><td>0</td><td>Colab clones the repo, then finds scripts/</td></tr>
    <tr><td>1</td><td>PDFs to speaker turns to Q&amp;A pairs</td></tr>
    <tr><td>2–3</td><td>BERTopic, then FinBERT on question and answer</td></tr>
    <tr><td>6</td><td>Three baselines, each with n</td></tr>
    <tr><td>8–9</td><td>Rules, then the PRA note</td></tr>
    </table>
  </div>
  <div>
    <h3>For a technical listener</h3>
    <p>State lives in data/boe.sqlite. CSV copies are in data/processed/. Human labels are scored, not rewritten.</p>
    <div class="callout info"><div class="t">If a live re-fit is too slow</div>
    Walk the saved outputs already in the notebook, then the PRA note for HSBC 2025-H1. Do not refit BERTopic on stage.</div>
  </div>
</div>
""",
        11))
    slides.append(slide(
        "Close · Speaks: Taz",
        "Sentiment alone is not enough",
        """
<div class="callout ok"><div class="t">Close</div>
The strongest result is a negative one. FinBERT and the lexicon agree about 58% of the time. FinBERT calls about 80% of questions neutral and misses the hedged negatives. Kappa is 0.28. We did not move the 70% bar. HSBC half-year 2025 is watch, not a rating. Credit Suisse’s failure quarter, held out of the pitch, is null.</div>
<div class="cols">
  <div class="callout info"><div class="t">For the Bank</div>
  Is alert / watch / null the right shape, and is a 45–90 minute re-run acceptable each quarter?</div>
  <div class="callout"><div class="t">For markers</div>
  Depth on two banks, an honest miss on the 70% bar, and a pipeline that can say null.</div>
</div>
""",
        12))
    # last footer should name the pdf
    slides[-1] = slides[-1].replace(
        "Group 9 · BoE employer project",
        "Group9_CAM_EP_Assignment3_presentation_slides.pdf",
    )
    html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"/>
<title>Group 9 · CAM EP Assignment 3</title>
<style>{SLIDE_CSS}</style></head><body>
{''.join(slides)}
</body></html>
"""
    path.write_text(html, encoding="utf-8")


def print_pdf(html: Path, pdf: Path) -> None:
    subprocess.check_call(
        [
            CHROME,
            "--headless",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf}",
            html.as_uri(),
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def main() -> None:
    n1, n2, n3 = words(S1), words(S2), words(S3)
    total = n1 + n2 + n3
    print(f"words section1={n1} section2={n2} section3={n3} total={total}")
    if not 1350 <= total <= 1650:
        raise SystemExit(f"word count {total} outside 1,500 ±10%")
    OUT.mkdir(parents=True, exist_ok=True)
    report_html = OUT / "Group9_CAM_EP_Assignment3_report.html"
    report_docx = OUT / "Group9_CAM_EP_Assignment3_report.docx"
    report_pdf = OUT / "Group9_CAM_EP_Assignment3_report.pdf"
    slides_html = OUT / "Group9_CAM_EP_Assignment3_presentation_slides.html"
    slides_pdf = OUT / "Group9_CAM_EP_Assignment3_presentation_slides.pdf"
    build_report_html(report_html)
    build_docx(report_docx)
    build_slides_html(slides_html)
    print_pdf(report_html, report_pdf)
    print_pdf(slides_html, slides_pdf)
    print("wrote", report_docx.name, report_pdf.name, slides_pdf.name)


if __name__ == "__main__":
    main()
